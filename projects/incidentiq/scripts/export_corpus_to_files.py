"""Export selected corpus records to multi-format files under data/documents/.

Run once (or after corpus edits) to regenerate TXT, DOCX, PDF, and XLSX samples:

    python scripts/export_corpus_to_files.py
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from data.sample_incidents import INCIDENTS, SOPS  # noqa: E402

OUTPUT_DIR = _PROJECT_ROOT / "data" / "documents"


def _format_numbered_steps(steps: list[str]) -> str:
    return " ".join(f"{i}) {step}" for i, step in enumerate(steps, start=1))


def format_incident_chunk(incident: dict[str, Any]) -> str:
    parts = [
        f"Incident {incident['id']} [{incident.get('severity', 'N/A')} - {incident.get('category', 'Unknown')}]: {incident['title']}.",
        f"Description: {incident.get('description', '').strip()}.",
        f"Triage Steps: {_format_numbered_steps(incident.get('triage_steps', []))}.",
        f"Root Cause: {incident.get('root_cause', '').strip()}.",
        f"Resolution: {_format_numbered_steps(incident.get('resolution_steps', []))}.",
        f"MTTR: {incident.get('mttr_minutes', 'N/A')} minutes.",
        f"Lessons Learned: {incident.get('lessons_learned', '').strip()}.",
    ]
    return " ".join(part for part in parts if part and part.strip() not in {".", "Description: ."})


def format_sop_chunk(sop: dict[str, Any]) -> str:
    parts = [
        f"SOP {sop['id']} [{sop['title']}] Version {sop.get('version', '1.0')}.",
        f"Applies to: {sop.get('applicability', '').strip()}.",
        f"Severity Trigger: {sop.get('severity_trigger', 'N/A')}.",
        f"Steps: {_format_numbered_steps(sop.get('steps', []))}.",
        f"Escalation: {sop.get('escalation_path', '').strip()}.",
        f"Owner: {sop.get('owner', 'Unassigned')}.",
    ]
    return " ".join(part for part in parts if part and part.strip() not in {".", "Applies to: ."})


def _find_incident(incident_id: str) -> dict:
    for inc in INCIDENTS:
        if inc["id"] == incident_id:
            return inc
    raise KeyError(incident_id)


def _find_sop(sop_id: str) -> dict:
    for sop in SOPS:
        if sop["id"] == sop_id:
            return sop
    raise KeyError(sop_id)


def export_txt(incident_id: str, filename: str) -> None:
    text = format_incident_chunk(_find_incident(incident_id))
    (OUTPUT_DIR / filename).write_text(text + "\n", encoding="utf-8")


def export_docx(sop_id: str, filename: str) -> None:
    from docx import Document

    text = format_sop_chunk(_find_sop(sop_id))
    doc = Document()
    for paragraph in text.split(". "):
        chunk = paragraph.strip()
        if chunk:
            doc.add_paragraph(chunk if chunk.endswith(".") else chunk + ".")
    doc.save(OUTPUT_DIR / filename)


def export_pdf(incident_id: str, filename: str) -> None:
    from fpdf import FPDF

    text = format_incident_chunk(_find_incident(incident_id))
    safe_text = (
        text.replace("\u2014", "-")
        .replace("\u2013", "-")
        .encode("ascii", errors="replace")
        .decode("ascii")
    )
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=9)
    pdf.write(5, safe_text)
    pdf.output(str(OUTPUT_DIR / filename))


def export_xlsx(filename: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Incidents"
    ws.append(["id", "title", "severity", "category", "description"])
    for inc in INCIDENTS[:12]:
        ws.append(
            [
                inc["id"],
                inc["title"],
                inc["severity"],
                inc["category"],
                inc.get("description", "")[:500],
            ]
        )
    wb.save(OUTPUT_DIR / filename)


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    export_txt("INC-001", "INC-001_postgres_pool_exhaustion.txt")
    export_docx("SOP-MQ-001", "SOP-MQ-001_kafka_consumer_lag.docx")
    export_pdf("INC-006", "INC-006_pod_crashloop.pdf")
    export_xlsx("incident_register.xlsx")
    print(f"Exported sample documents to {OUTPUT_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
